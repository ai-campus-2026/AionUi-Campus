"""文档定位与文本提取

把一个"文件名字符串"解析为可用于对比的文档对象，解析顺序：
  1. 本地路径（绝对路径 / 相对 CWD 的路径）
  2. policy-search 知识库标题 / doc_id / 原始文件名（"规则库"场景）
  3. 配置的搜索目录内按文件名模糊匹配（"上传文件"场景，目录见 POLICY_COMPARE_SEARCH_DIRS）

支持格式：.pdf / .docx / .txt / .md / .json（知识库结构化 JSON，取 raw_text）
"""

import json
import logging
import os
import re
import unicodedata
from difflib import SequenceMatcher
from typing import Any, Dict, List, Optional, Tuple

from config import Config

logger = logging.getLogger(__name__)

# 名称匹配时忽略的标点/修饰字符
_PUNCT_PATTERN = re.compile(r"[\s《》〈〉“”\"'‘’（）()【】\[\]、，。,.：:；;！!？?—－\-_·/\\|~]+")


class DocumentNotFound(Exception):
    """未能定位到政策文档（携带候选列表，供 Agent 参考）"""

    def __init__(self, spec: str, candidates: Optional[List[str]] = None):
        super().__init__(f"未找到政策文件: {spec}")
        self.spec = spec
        self.candidates = candidates or []


class DocumentError(Exception):
    """文件存在但无法读取/解析"""


# ============================================================
# 文本清洗与提取
# ============================================================


def sanitize_text(text: str) -> str:
    """清理文本，确保输出为合法 UTF-8（与 policy-search 同策略，轻量版）"""
    if not text:
        return ""
    text = unicodedata.normalize("NFC", text)
    replacements = {
        "\u2212": "-",
        "\u2013": "-",
        "\u2014": "--",
        "\u2018": "'",
        "\u2019": "'",
        "\u201c": '"',
        "\u201d": '"',
        "\u2026": "...",
        "\u00a0": " ",
        "\u200b": "",
        "\u200c": "",
        "\u200d": "",
        "\ufeff": "",
        "\u00ad": "",
        "\ufffd": "?",
    }
    for old, new in replacements.items():
        text = text.replace(old, new)
    text = "".join(ch for ch in text if ch in ("\n", "\r", "\t") or unicodedata.category(ch)[0] != "C")
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text


# 抽取文本"乱码率"阈值：控制字符（不含 \n\r\t）与 U+FFFD 替换符占比超过此值判定为不可读。
# 实测：PyPDF2 不支持 /UniGB-UCS2-H 等中文 CMap 时乱码率约 0.24，pdfplumber 正常抽取约 0.00。
_GARBAGE_OK_THRESHOLD = 0.05


def _garbage_ratio(text: str) -> float:
    """估算文本乱码率：控制字符（不含换行/制表）与 U+FFFD 替换符占比。

    用于识别 PyPDF2 遇到不支持的中文 CMap（如 /UniGB-UCS2-H）时产出的乱码——
    这类乱码以大量 C1 控制字符为特征，而正常中文文本该比例接近 0。
    """
    if not text:
        return 1.0
    sample = text[:8000]
    bad = 0
    for ch in sample:
        if ch == "\ufffd":
            bad += 1
            continue
        if unicodedata.category(ch)[0] == "C" and ch not in ("\n", "\r", "\t"):
            bad += 1
    return bad / max(1, len(sample))


def _read_text_file(file_path: str) -> str:
    """尝试多种编码读取文本文件"""
    encodings = ["utf-8", "utf-8-sig", "gbk", "gb2312", "gb18030", "big5", "latin-1"]
    for enc in encodings:
        try:
            with open(file_path, "r", encoding=enc) as f:
                return f.read()
        except (UnicodeDecodeError, UnicodeError):
            continue
    with open(file_path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def _extract_pdf_text(pdf_path: str) -> str:
    """PDF 文本提取：PyPDF2 优先，乱码则自动回退 pdfplumber，最终选乱码更少的一方。

    背景：PyPDF2 不支持部分中文 CMap（如 /UniGB-UCS2-H），对这类 PDF 会抽出乱码
    （数字/ASCII 正常、中文成控制字符乱码），但长度仍 >100，旧逻辑会直接返回乱码。
    现改为按"乱码率"择优：PyPDF2 结果可读时走快路径，否则回退 pdfplumber。
    """
    pypdf2_text: Optional[str] = None
    try:
        from PyPDF2 import PdfReader

        parts = []
        reader = PdfReader(pdf_path)
        for page in reader.pages:
            page_text = page.extract_text() or ""
            if page_text.strip():
                parts.append(page_text)
        pypdf2_text = "\n".join(parts)
    except ImportError:
        logger.warning("PyPDF2 未安装，跳过")
    except Exception as e:
        logger.warning(f"PyPDF2 提取失败: {e}")

    # 快路径：PyPDF2 抽取结果可读且足够长，直接返回（避免对正常 PDF 多跑一次 pdfplumber）
    if (
        pypdf2_text
        and pypdf2_text.strip()
        and len(pypdf2_text.strip()) > 100
        and _garbage_ratio(pypdf2_text) < _GARBAGE_OK_THRESHOLD
    ):
        logger.info(f"PyPDF2 提取成功，文本长度: {len(pypdf2_text)}")
        return pypdf2_text

    pdfplumber_text: Optional[str] = None
    try:
        import pdfplumber

        parts = []
        with pdfplumber.open(pdf_path) as pdf:
            for page in pdf.pages:
                page_text = page.extract_text() or ""
                if page_text.strip():
                    parts.append(page_text)
        pdfplumber_text = "\n".join(parts)
    except ImportError:
        logger.warning("pdfplumber 未安装，跳过")
    except Exception as e:
        logger.warning(f"pdfplumber 提取失败: {e}")

    candidates = [
        (name, text)
        for name, text in (("PyPDF2", pypdf2_text), ("pdfplumber", pdfplumber_text))
        if text and text.strip()
    ]
    if not candidates:
        raise DocumentError(f"PDF 文本提取失败（可能为扫描件或加密文件）: {os.path.basename(pdf_path)}")

    # 选乱码率最低的一份；并列时取更长的一份
    best_name, best_text = min(candidates, key=lambda nt: (_garbage_ratio(nt[1]), -len(nt[1])))
    if _garbage_ratio(best_text) >= _GARBAGE_OK_THRESHOLD:
        logger.warning(
            f"两种方式抽取均疑似乱码（选 {best_name}），结果可能含乱码: {os.path.basename(pdf_path)}"
        )
    else:
        logger.info(f"{best_name} 提取成功（已回退乱码来源），文本长度: {len(best_text)}")
    return best_text


def _extract_docx_text(docx_path: str) -> str:
    """DOCX 文本提取（python-docx，按段落顺序）"""
    try:
        import docx

        document = docx.Document(docx_path)
        parts = [p.text for p in document.paragraphs if p.text and p.text.strip()]
        # 表格内文字也一并提取
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text and cell.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        text = "\n".join(parts)
        logger.info(f"python-docx 提取成功，文本长度: {len(text)}")
        return text
    except ImportError:
        raise DocumentError("python-docx 未安装，无法解析 .docx 文件。请运行: pip install python-docx")
    except Exception as e:
        raise DocumentError(f"DOCX 解析失败: {e}")


def _load_json_document(file_path: str) -> Tuple[str, str, str]:
    """读取知识库格式的结构化 JSON：返回 (text, name, version)"""
    try:
        with open(file_path, "r", encoding="utf-8") as f:
            data = json.load(f)
    except Exception as e:
        raise DocumentError(f"JSON 解析失败: {e}")

    meta = data.get("meta") or {}
    raw_text = (data.get("raw_text") or "").strip()
    if not raw_text:
        raw_text = _reconstruct_from_requirements(data)
    name = meta.get("title") or os.path.splitext(os.path.basename(file_path))[0]
    version = _version_from_meta(meta, name)
    if not raw_text:
        raise DocumentError(f"JSON 中没有可对比的文本内容: {os.path.basename(file_path)}")
    return raw_text, name, version


def _reconstruct_from_requirements(data: Dict[str, Any]) -> str:
    """兜底：从结构化 requirements 的 source_quote 重建文本"""
    lines = []
    for _, section in (data.get("requirements") or {}).items():
        for cond in section.get("conditions") or []:
            quote = (cond.get("source_quote") or "").strip()
            item = (cond.get("item") or "").strip()
            if quote:
                lines.append(quote)
            elif item:
                lines.append(item)
    return "\n".join(lines)


def load_text_from_file(file_path: str) -> Tuple[str, str, str]:
    """按扩展名提取文本：返回 (text, name, version)；version 可能为空字符串"""
    ext = os.path.splitext(file_path)[1].lower()
    stem = os.path.splitext(os.path.basename(file_path))[0]

    if ext == ".pdf":
        text = sanitize_text(_extract_pdf_text(file_path))
    elif ext in (".txt", ".md", ".text"):
        text = sanitize_text(_read_text_file(file_path))
    elif ext == ".docx":
        text = sanitize_text(_extract_docx_text(file_path))
    elif ext == ".doc":
        raise DocumentError(
            "不支持 .doc 格式（旧版 Word 二进制格式），请将文件另存为 .docx 后重试。"
            "（Word 中：文件 → 另存为 → 选择 .docx 格式）"
        )
    elif ext == ".json":
        text, name, version = _load_json_document(file_path)
        return sanitize_text(text), name, version
    else:
        raise DocumentError(f"不支持的文件格式: {ext or '(无扩展名)'}，仅支持 .pdf / .docx / .txt / .md / .json")

    if not text.strip():
        raise DocumentError(f"文件解析后内容为空: {os.path.basename(file_path)}")
    return text, stem, infer_version(stem)


# ============================================================
# 名称归一化 / 相似度
# ============================================================


def _normalize(text: str) -> str:
    """归一化：NFKC 全角转半角、去空白标点、转小写"""
    if not text:
        return ""
    text = unicodedata.normalize("NFKC", text)
    text = _PUNCT_PATTERN.sub("", text)
    return text.lower()


def _score(query: str, target: str) -> float:
    """名称相似度打分（0~1）"""
    q, t = _normalize(query), _normalize(target)
    if not q or not t:
        return 0.0
    if q == t:
        return 1.0
    if q in t or t in q:
        return 0.92
    return SequenceMatcher(None, q, t).ratio()


def _file_variants(filename: str) -> List[str]:
    """文件名及其去扩展名两种写法"""
    stem = os.path.splitext(filename)[0]
    return [filename, stem] if stem != filename else [filename]


def infer_version(name: str) -> str:
    """从名称中推断版本号，如 'xxx办法2025.pdf' → '2025版'"""
    if not name:
        return ""
    m = re.search(r"(20\d{2})", name)
    return f"{m.group(1)}版" if m else ""


def _version_from_meta(meta: Dict[str, Any], fallback_name: str) -> str:
    year = meta.get("year")
    if year:
        return f"{year}版"
    version = infer_version(meta.get("title") or "")
    if version:
        return version
    date = str(meta.get("effective_date") or "")
    m = re.search(r"(20\d{2})", date)
    return f"{m.group(1)}版" if m else ""


# ============================================================
# 知识库 / 搜索目录
# ============================================================


def _load_kb_docs() -> List[Dict[str, Any]]:
    """加载知识库索引 +（必要时）详情中的 meta，返回可匹配的文档列表"""
    index_path = os.path.join(Config.KNOWLEDGE_BASE_DIR, "index.json")
    if not os.path.isfile(index_path):
        logger.info(f"知识库索引不存在，跳过 KB 匹配: {index_path}")
        return []

    try:
        with open(index_path, "r", encoding="utf-8") as f:
            index = json.load(f)
    except Exception as e:
        logger.warning(f"知识库索引读取失败: {e}")
        return []

    docs: List[Dict[str, Any]] = []
    for category, entries in (index.get("categories") or {}).items():
        for entry in entries or []:
            doc_id = entry.get("doc_id") or ""
            detail_path = os.path.join(Config.KNOWLEDGE_BASE_DIR, category, f"{doc_id}.json")
            title = entry.get("title") or doc_id
            source_file = ""
            year = entry.get("year")
            if os.path.isfile(detail_path):
                try:
                    with open(detail_path, "r", encoding="utf-8") as f:
                        meta = (json.load(f).get("meta") or {})
                    title = meta.get("title") or title
                    source_file = meta.get("source_file") or ""
                    year = meta.get("year") or year
                except Exception:
                    pass
            docs.append(
                {
                    "doc_id": doc_id,
                    "category": category,
                    "title": title,
                    "source_file": os.path.basename(source_file) if source_file else "",
                    "year": year,
                    "detail_path": detail_path,
                }
            )
            if len(docs) >= 50:
                return docs
    return docs


def _match_kb(spec: str) -> Optional[Dict[str, Any]]:
    """在知识库中按标题 / doc_id / 原始文件名匹配"""
    best = None
    for doc in _load_kb_docs():
        targets = [doc["title"], doc["doc_id"]] + _file_variants(doc["source_file"] or "")
        score = max(_score(spec, t) for t in targets if t)
        if score >= Config.MATCH_THRESHOLD and (best is None or score > best["score"]):
            best = {"score": score, "doc": doc}
    return best


def _iter_search_files():
    """遍历搜索目录（受限深度与文件数），产出候选文件路径"""
    visited = 0
    for root in Config.SEARCH_DIRS:
        root_depth = root.rstrip("\\/").count(os.sep)
        for dirpath, dirnames, filenames in os.walk(root):
            depth = dirpath.count(os.sep) - root_depth
            dirnames[:] = [d for d in dirnames if not d.startswith(".")]
            if depth >= Config.SEARCH_MAX_DEPTH:
                dirnames[:] = []
            for filename in filenames:
                if not filename.lower().endswith(Config.SUPPORTED_EXTS):
                    continue
                visited += 1
                if visited > Config.SEARCH_MAX_FILES:
                    logger.warning(f"搜索目录文件数超过上限 {Config.SEARCH_MAX_FILES}，提前结束")
                    return
                yield os.path.join(dirpath, filename)


def _match_search_dirs(spec: str) -> Optional[Dict[str, Any]]:
    """在搜索目录中按文件名匹配"""
    best = None
    for file_path in _iter_search_files():
        basename = os.path.basename(file_path)
        score = max(_score(spec, v) for v in _file_variants(basename))
        if score >= Config.MATCH_THRESHOLD and (best is None or score > best["score"]):
            best = {"score": score, "path": file_path}
    return best


def _recent_search_files(limit: int) -> List[str]:
    """搜索目录中最近修改的文件（mtime 倒序，供候选提示）"""
    candidates: List[Tuple[float, str]] = []
    for path in _iter_search_files():
        try:
            candidates.append((os.path.getmtime(path), path))
        except OSError:
            continue
    candidates.sort(key=lambda x: x[0], reverse=True)
    return [p for _, p in candidates[:limit]]


def _collect_candidates(spec: str, limit: int = 8) -> List[str]:
    """收集候选（知识库标题 + 搜索目录中相似文件名），用于报错提示"""
    scored: List[Tuple[float, str]] = []
    for doc in _load_kb_docs():
        score = max(
            _score(spec, doc["title"]),
            _score(spec, doc["doc_id"]),
            max((_score(spec, v) for v in _file_variants(doc["source_file"] or "")), default=0.0),
        )
        if score >= 0.3:
            scored.append((score, f"{doc['title']}（规则库·{doc['category']}）"))
    for file_path in _iter_search_files():
        basename = os.path.basename(file_path)
        score = max(_score(spec, v) for v in _file_variants(basename))
        if score >= 0.3:
            scored.append((score, f"{basename}（{os.path.dirname(file_path)}）"))

    # 相似项太少时补充“可用文档”，帮助用户确认正确名称
    if len(scored) < 3:
        for doc in _load_kb_docs():
            scored.append((0.0, f"{doc['title']}（规则库·{doc['category']}）"))
        for file_path in _recent_search_files(5):
            basename = os.path.basename(file_path)
            scored.append((0.0, f"{basename}（{os.path.dirname(file_path)}）"))

    # 稳定排序：同分保持原有优先级（相似项在前，补充项在后）
    scored.sort(key=lambda x: x[0], reverse=True)
    # 去重后取前 N 个
    seen, result = set(), []
    for _, label in scored:
        if label not in seen:
            seen.add(label)
            result.append(label)
        if len(result) >= limit:
            break
    return result


# ============================================================
# 对外入口
# ============================================================


def resolve_document(spec: str) -> Dict[str, Any]:
    """
    将输入字符串解析为文档对象。

    Returns:
        {
          "name": 文档名, "version": 版本号（可能为空）,
          "text": 全文, "origin": "path"|"kb"|"search",
          "path"/"doc_id": 来源标识
        }

    Raises:
        DocumentNotFound: 三条路都无法定位（含候选列表）
        DocumentError: 定位成功但读取/解析失败
    """
    spec = (spec or "").strip().strip("《》\"'")
    if not spec:
        raise DocumentError("文件名/路径不能为空")

    # 1) 直接路径
    path = os.path.abspath(os.path.expanduser(spec))
    if os.path.isfile(path):
        text, name, version = load_text_from_file(path)
        logger.info(f"[resolve] 路径命中: {path}")
        return {"name": name, "version": version, "text": text, "origin": "path", "path": path}

    # 2) 知识库 / 3) 搜索目录
    kb_hit = _match_kb(spec)
    dir_hit = _match_search_dirs(spec)

    pick: Optional[Dict[str, Any]] = None
    if kb_hit and dir_hit:
        pick = kb_hit if kb_hit["score"] >= dir_hit["score"] else dir_hit
    else:
        pick = kb_hit or dir_hit

    if pick is None:
        raise DocumentNotFound(spec, _collect_candidates(spec))

    if "doc" in pick:  # 知识库命中
        doc = pick["doc"]
        logger.info(f"[resolve] 知识库命中: {doc['title']} (score={pick['score']:.2f})")
        text, name, version = _load_json_document(doc["detail_path"])
        return {
            "name": name,
            "version": version,
            "text": text,
            "origin": "kb",
            "doc_id": doc["doc_id"],
            "path": doc["detail_path"],
        }

    # 搜索目录命中
    file_path = pick["path"]
    logger.info(f"[resolve] 搜索目录命中: {file_path} (score={pick['score']:.2f})")
    text, name, version = load_text_from_file(file_path)
    return {"name": name, "version": version, "text": text, "origin": "search", "path": file_path}
